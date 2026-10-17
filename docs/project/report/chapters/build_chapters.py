import re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = sys.argv[1]
OUT = sys.argv[2]
REFS_OUT = sys.argv[3]
AR, EN = "Simplified Arabic", "Times New Roman"

ACC = "accessed Oct. 13, 2026"
REFS = {
 "who_diet": f'World Health Organization, "Healthy diet," Fact sheet, Apr. 2020. [Online]. Available: https://www.who.int/news-room/fact-sheets/detail/healthy-diet ({ACC}).',
 "who_ncd": f'World Health Organization, "Noncommunicable diseases," Fact sheet. [Online]. Available: https://www.who.int/news-room/fact-sheets/detail/noncommunicable-diseases ({ACC}).',
 "ada2025": 'American Diabetes Association Professional Practice Committee, "5. Facilitating positive health behaviors and well-being to improve health outcomes: Standards of Care in Diabetes—2025," Diabetes Care, vol. 48, suppl. 1, 2025.',
 "appel1997": 'L. J. Appel et al., "A clinical trial of the effects of dietary patterns on blood pressure," N. Engl. J. Med., vol. 336, no. 16, pp. 1117–1124, 1997.',
 "aha2021": 'A. H. Lichtenstein et al., "2021 dietary guidance to improve cardiovascular health: A scientific statement from the American Heart Association," Circulation, vol. 144, no. 23, pp. e472–e487, 2021.',
 "trattner2017": 'C. Trattner and D. Elsweiler, "Food recommender systems: Important contributions, challenges and future research directions," arXiv:1711.02760, 2017.',
 "min2019": 'W. Min, S. Jiang, L. Liu, Y. Rui, and R. Jain, "A survey on food computing," ACM Comput. Surv., vol. 52, no. 5, pp. 1–36, 2019, doi: 10.1145/3329168.',
 "bawajeeh2021": 'A. O. Bawajeeh et al., "Development of an Arabic food composition database for use in an Arabic online dietary assessment tool (myfood24)," J. Food Compos. Anal., vol. 102, Art. no. 104047, 2021, doi: 10.1016/j.jfca.2021.104047.',
 "fda2019": 'U.S. Food and Drug Administration, "General wellness: Policy for low risk devices," Guidance for Industry and FDA Staff, Sep. 2019.',
 "fdc": f'U.S. Department of Agriculture, Agricultural Research Service, "FoodData Central." [Online]. Available: https://fdc.nal.usda.gov ({ACC}).',
 "fndds": f'U.S. Department of Agriculture, Agricultural Research Service, "Food and Nutrient Database for Dietary Studies (FNDDS)." [Online]. Available: https://www.ars.usda.gov/northeast-area/beltsville-md-bhnrc/beltsville-human-nutrition-research-center/food-surveys-research-group/docs/fndds/ ({ACC}).',
 "ortools": f'L. Perron and V. Furnon, "OR-Tools," Google. [Online]. Available: https://developers.google.com/optimization ({ACC}).',
 "stigler1945": 'G. J. Stigler, "The cost of subsistence," J. Farm Econ., vol. 27, no. 2, pp. 303–314, 1945.',
 "dantzig1990": 'G. B. Dantzig, "The diet problem," Interfaces, vol. 20, no. 4, pp. 43–47, 1990, doi: 10.1287/inte.20.4.43.',
 "who_sodium2012": 'World Health Organization, Guideline: Sodium Intake for Adults and Children. Geneva, Switzerland: WHO, 2012.',
 "who_satfat2023": 'World Health Organization, Saturated Fatty Acid and Trans-Fatty Acid Intake for Adults and Children: WHO Guideline. Geneva, Switzerland: WHO, 2023.',
 "harris1918": 'J. A. Harris and F. G. Benedict, "A biometric study of human basal metabolism," Proc. Natl. Acad. Sci. USA, vol. 4, no. 12, pp. 370–373, 1918.',
 "mifflin1990": 'M. D. Mifflin, S. T. St Jeor, L. A. Hill, B. J. Scott, S. A. Daugherty, and Y. O. Koh, "A new predictive equation for resting energy expenditure in healthy individuals," Am. J. Clin. Nutr., vol. 51, no. 2, pp. 241–247, 1990.',
 "iom2005": 'Institute of Medicine, Dietary Reference Intakes for Energy, Carbohydrate, Fiber, Fat, Fatty Acids, Cholesterol, Protein, and Amino Acids. Washington, DC, USA: National Academies Press, 2005.',
 "efsa2010": 'EFSA Panel on Dietetic Products, Nutrition, and Allergies, "Scientific opinion on dietary reference values for water," EFSA J., vol. 8, no. 3, Art. no. 1459, 2010.',
 "seljak2009": 'B. Koroušić Seljak, "Computer-based dietary menu planning," J. Food Compos. Anal., vol. 22, no. 5, pp. 414–420, 2009.',
 "burke2002": 'R. Burke, "Hybrid recommender systems: Survey and experiments," User Model. User-Adapt. Interact., vol. 12, no. 4, pp. 331–370, 2002.',
 "ricci2022": 'F. Ricci, L. Rokach, and B. Shapira, Eds., Recommender Systems Handbook, 3rd ed. New York, NY, USA: Springer, 2022.',
 "rossi2006": 'F. Rossi, P. van Beek, and T. Walsh, Eds., Handbook of Constraint Programming. Amsterdam, The Netherlands: Elsevier, 2006.',
 "elsweiler2017": 'D. Elsweiler, C. Trattner, and M. Harvey, "Exploiting food choice biases for healthier recipe recommendation," in Proc. 40th Int. ACM SIGIR Conf. Res. Dev. Inf. Retr., 2017, doi: 10.1145/3077136.3080826.',
 "tintarev2012": 'N. Tintarev and J. Masthoff, "Evaluating the effectiveness of explanations for recommender systems," User Model. User-Adapt. Interact., vol. 22, no. 4–5, pp. 399–439, 2012, doi: 10.1007/s11257-011-9117-5.',
 "burke2011": 'L. E. Burke, J. Wang, and M. A. Sevick, "Self-monitoring in weight loss: A systematic review of the literature," J. Am. Diet. Assoc., vol. 111, no. 1, pp. 92–102, 2011.',
 "michie2013": 'S. Michie et al., "The behavior change technique taxonomy (v1) of 93 hierarchically clustered techniques," Ann. Behav. Med., vol. 46, no. 1, pp. 81–95, 2013.',
 "deterding2011": 'S. Deterding, D. Dixon, R. Khaled, and L. Nacke, "From game design elements to gamefulness: Defining \'gamification\'," in Proc. 15th Int. Academic MindTrek Conf., 2011, pp. 9–15.',
 "stonge2017": 'M.-P. St-Onge et al., "Meal timing and frequency: Implications for cardiovascular disease prevention: A scientific statement from the American Heart Association," Circulation, vol. 135, no. 9, pp. e96–e121, 2017.',
 "ge2015": 'M. Ge, F. Ricci, and D. Massimo, "Health-aware food recommender system," in Proc. 9th ACM Conf. Recommender Systems (RecSys \'15), 2015, doi: 10.1145/2792838.2796554.',
 "yera2019": 'R. Yera Toledo, A. A. Alzahrani, and L. Martínez, "A food recommender system considering nutritional information and user preferences," IEEE Access, vol. 7, pp. 96695–96711, 2019, doi: 10.1109/ACCESS.2019.2929413.',
}

text = open(SRC, encoding="utf-8").read()
order = []
for k in re.findall(r"\[@([a-z0-9_]+)\]", text):
    if k not in REFS: raise SystemExit("missing ref " + k)
    if k not in order: order.append(k)
num = {k: i + 1 for i, k in enumerate(order)}
text = re.sub(r"\[@([a-z0-9_]+)\]", lambda m: f"[{num[m.group(1)]}]", text)

doc = Document()
st = doc.styles["Normal"]
st.font.name = AR; st.font.size = Pt(14)
rpr = st.element.get_or_add_rPr()
rf = rpr.find(qn("w:rFonts"))
rf = rf if rf is not None else OxmlElement("w:rFonts")
for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), AR)
if rf.getparent() is None: rpr.append(rf)
szcs = OxmlElement("w:szCs"); szcs.set(qn("w:val"), "28"); rpr.append(szcs)
st.paragraph_format.line_spacing = 1.5
st.paragraph_format.space_after = Pt(6)
for name, size in (("Heading 1", 20), ("Heading 2", 16), ("Heading 3", 14)):
    h = doc.styles[name]
    h.font.name = AR; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = RGBColor(0, 0, 0)
    hr = h.element.get_or_add_rPr()
    f = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): f.set(qn(a), AR)
    hr.append(f)
    b = OxmlElement("w:bCs"); hr.append(b)
    s2 = OxmlElement("w:szCs"); s2.set(qn("w:val"), str(size * 2)); hr.append(s2)
    h.paragraph_format.space_before = Pt(12); h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True

def bidi(p, on=True):
    ppr = p._p.get_or_add_pPr()
    e = OxmlElement("w:bidi")
    if not on: e.set(qn("w:val"), "0")
    ppr.append(e)

def run(p, t, font=AR, size=14, bold=False, rtl=True, italic=False):
    r = p.add_run(t)
    r.font.name = font; r.font.size = Pt(size); r.bold = bold; r.italic = italic
    rp = r._r.get_or_add_rPr()
    f = rp.find(qn("w:rFonts"))
    if f is None: f = OxmlElement("w:rFonts"); rp.insert(0, f)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): f.set(qn(a), font)
    if bold: rp.append(OxmlElement("w:bCs"))
    s = OxmlElement("w:szCs"); s.set(qn("w:val"), str(int(size * 2))); rp.append(s)
    if rtl: rp.append(OxmlElement("w:rtl"))
    return r

def rich(p, t, size=14, rtl=True, font=AR):
    for i, part in enumerate(re.split(r"\*\*", t)):
        if part: run(p, part, font=font, size=size, bold=(i % 2 == 1), rtl=rtl)

def para(t="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=14, rtl=True, font=AR, style=None, sa=None):
    p = doc.add_paragraph(style=style)
    bidi(p, rtl)
    p.alignment = align
    if sa is not None: p.paragraph_format.space_after = Pt(sa)
    if t: rich(p, t, size, rtl, font)
    return p

def footer(sec):
    sec.footer.is_linked_to_previous = False
    p = sec.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for typ, txt in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
        r = p.add_run()
        if typ:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), typ); r._r.append(fc)
        elif txt == " PAGE ":
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._r.append(it)
        else:
            r.text = txt; r.font.size = Pt(12)

def section(arabic=True, fmt=None, start=None, first=False):
    sec = doc.sections[0] if first else doc.add_section(WD_SECTION.NEW_PAGE)
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    if arabic: sec.right_margin, sec.left_margin = Cm(3), Cm(2)
    else: sec.left_margin, sec.right_margin = Cm(3), Cm(2)
    sp = sec._sectPr
    for e in sp.findall(qn("w:pgNumType")): sp.remove(e)
    pn = OxmlElement("w:pgNumType")
    if fmt: pn.set(qn("w:fmt"), fmt)
    if start is not None: pn.set(qn("w:start"), str(start))
    sp.append(pn)
    if arabic:
        b = OxmlElement("w:bidi"); sp.append(b)
    footer(sec)
    return sec

def heading(t, level):
    h = doc.add_heading(level=level); bidi(h)
    h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(h, t, size={1: 20, 2: 16, 3: 14}[level], bold=True)
    return h

def table(caption, rows):
    cp = para(align=WD_ALIGN_PARAGRAPH.CENTER, sa=4)
    run(cp, "جدول " + caption, size=12, bold=True)
    t = doc.add_table(rows=0, cols=len(rows[0])); t.style = "Table Grid"; t.autofit = False
    W = {"1.1": [1.6, 7.4, 4.4, 2.6], "2.1": [2.8, 4.2, 4.5, 4.5], "2.2": [3.4, 4.0, 4.0, 4.6]}[caption.split(":")[0]]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tp = t._tbl.tblPr; bv = OxmlElement("w:bidiVisual"); tp.append(bv)
    for i, r in enumerate(rows):
        cells = t.add_row().cells
        for j, (c, v) in enumerate(zip(cells, r)):
            c.width = Cm(W[j])
            p = c.paragraphs[0]; bidi(p); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(2)
            rich(p, v, size=11)
            if i == 0:
                for r_ in p.runs: r_.bold = True
                sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), "D9E2F3")
                c._tc.get_or_add_tcPr().append(sh)
    para("", sa=2)

def process(body):
    for ln in body.split("\n"):
        if ln.startswith("### "): heading(ln[4:], 3)
        elif ln.startswith("## "): heading(ln[3:], 2)
        elif ln.startswith("- "): para(ln[2:], style="List Bullet")
        elif re.match(r"^\s{3}- ", ln): para(ln.strip()[2:], style="List Bullet 2")
        elif re.match(r"^\d+\. ", ln):
            p = para(ln, sa=4); p.paragraph_format.right_indent = Cm(0.5)
        elif ln.strip(): para(ln.strip())

ORD = {1: "الأول", 2: "الثاني", 3: "الثالث", 4: "الرابع", 5: "الخامس"}
blocks = re.split(r"^@@", text, flags=re.M)
first_chapter = True
started = False
for blk in blocks:
    if not blk.strip(): continue
    head, _, body = blk.partition("\n")
    if head == "ABSTRACT_AR":
        section(arabic=True, fmt="upperRoman", start=2, first=not started); started = True
        heading("الملخص", 1)
        for ln in body.strip().split("\n"):
            if ln.strip(): para(ln.strip())
    elif head == "ABSTRACT_EN":
        section(arabic=False, fmt="upperRoman")
        h = doc.add_heading(level=1); bidi(h, False); h.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run(h, "Abstract", font=EN, size=18, bold=True, rtl=False)
        for ln in body.strip().split("\n"):
            if ln.strip(): para(ln.strip(), rtl=False, font=EN)
    elif head.startswith("CHAPTER"):
        n, title = head.split(" ", 1)[1].split("|")
        n = int(n)
        section(arabic=True, fmt="decimal", start=1 if first_chapter else None); first_chapter = False
        for _ in range(8): para("", sa=12)
        h = doc.add_heading(level=1); bidi(h); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(h, "الفصل " + ORD[n], size=28, bold=True)
        h.add_run().add_break(WD_BREAK.LINE)
        run(h, title, size=28, bold=True)
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        process(body)
    elif head.startswith("TABLE"):
        cap = head.split(" ", 1)[1].replace("|", ": ", 1)
        rows = [r.split("|") for r in body.strip().split("\n") if r.strip()]
        table(cap, rows)
    elif head.startswith("END"):
        process(body)

# references
section(arabic=True)
heading("قائمة المراجع", 1)
for k in order:
    p = para(f"[{num[k]}] {REFS[k]}", align=WD_ALIGN_PARAGRAPH.LEFT, size=12, rtl=False, font=EN, sa=6)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(0.9); p.paragraph_format.first_line_indent = Cm(-0.9)
doc.save(OUT)

with open(REFS_OUT, "w", encoding="utf-8") as f:
    for k in order: f.write(f"[{num[k]}] {REFS[k]}\n")
print(len(order), "refs")
